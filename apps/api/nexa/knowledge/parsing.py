"""Document parsers: PDF, DOCX, TXT, website URL, FAQ and manual text."""

from __future__ import annotations

import io
import json
import re
from html.parser import HTMLParser


class ParseError(ValueError):
    pass


class _TextExtractor(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "nav", "footer", "header", "form"}
    BLOCK = {"p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6", "br", "tr", "section", "article", "td"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skip_depth = 0
        self.title = ""
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip_depth += 1
        if tag == "title":
            self._in_title = True
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip_depth:
            self.skip_depth -= 1
        if tag == "title":
            self._in_title = False
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self.skip_depth:
            self.parts.append(data)


def html_to_text(html: str) -> tuple[str, str]:
    p = _TextExtractor()
    p.feed(html)
    text = re.sub(r"[ \t]+", " ", "".join(p.parts))
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return p.title.strip(), text.strip()


def parse_document(source_type: str, raw: bytes) -> list[dict]:
    """Return sections: [{"text": ..., "meta": {...}}]."""
    if source_type in ("txt", "text"):
        return [{"text": raw.decode("utf-8", errors="replace"), "meta": {}}]
    if source_type == "faq":
        try:
            items = json.loads(raw)
        except json.JSONDecodeError as exc:
            raise ParseError("The FAQ must be a list of questions and answers.") from exc
        sections = []
        for i, item in enumerate(items):
            q, a = (item.get("question") or "").strip(), (item.get("answer") or "").strip()
            if q and a:
                sections.append({"text": f"{q}\n{a}", "meta": {"faq_index": i, "question": q, "atomic": True}})
        if not sections:
            raise ParseError("The FAQ has no complete question/answer pairs.")
        return sections
    if source_type == "pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(raw))
        except Exception as exc:  # noqa: BLE001 - pypdf raises many types
            raise ParseError("The PDF file could not be read.") from exc
        pages = [{"text": page.extract_text() or "", "meta": {"page": i + 1}} for i, page in enumerate(reader.pages)]
        pages = [p for p in pages if p["text"].strip()]
        if not pages:
            raise ParseError("The PDF has no readable text (scanned PDFs need OCR).")
        return pages
    if source_type == "docx":
        import docx

        try:
            d = docx.Document(io.BytesIO(raw))
        except Exception as exc:  # noqa: BLE001
            raise ParseError("The Word document could not be read.") from exc
        text = "\n\n".join(p.text for p in d.paragraphs if p.text.strip())
        for tbl in d.tables:
            for row in tbl.rows:
                text += "\n" + " | ".join(c.text.strip() for c in row.cells)
        return [{"text": text, "meta": {}}]
    if source_type == "url":
        title, text = html_to_text(raw.decode("utf-8", errors="replace"))
        return [{"text": text, "meta": {"page_title": title}}]
    raise ParseError(f"Unsupported document type: {source_type}")


def chunk_sections(sections: list[dict], max_chars: int = 900, overlap: int = 120) -> list[dict]:
    chunks: list[dict] = []
    for sec in sections:
        text = sec["text"].strip()
        if not text:
            continue
        if sec["meta"].get("atomic") or len(text) <= max_chars:
            chunks.append({"content": text, "meta": sec["meta"]})
            continue
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
        buf = ""
        for para in paragraphs:
            while len(para) > max_chars:  # very long paragraph: split on sentence boundaries
                cut = max(para.rfind(".", 0, max_chars), para.rfind("。", 0, max_chars), para.rfind("؟", 0, max_chars))
                cut = cut + 1 if cut > max_chars // 2 else max_chars
                piece, para = para[:cut], para[cut:].strip()
                if buf:
                    chunks.append({"content": buf, "meta": sec["meta"]})
                    buf = ""
                chunks.append({"content": piece.strip(), "meta": sec["meta"]})
            if len(buf) + len(para) + 1 > max_chars and buf:
                chunks.append({"content": buf, "meta": sec["meta"]})
                buf = buf[-overlap:].split(" ", 1)[-1] + "\n" + para if overlap else para
            else:
                buf = f"{buf}\n{para}" if buf else para
        if buf:
            chunks.append({"content": buf, "meta": sec["meta"]})
    return chunks
